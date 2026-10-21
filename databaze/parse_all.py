import os
import re
import glob
import json
import xml.etree.ElementTree as ET
from pathlib import Path
import pandas as pd
import duckdb
import yaml

try:
    from yaml import CSafeLoader as YamlLoader
except ImportError:
    from yaml import SafeLoader as YamlLoader

try:
    from python_calamine import CalamineWorkbook
    HAS_CALAMINE = True
except ImportError:
    import openpyxl
    CalamineWorkbook = None
    HAS_CALAMINE = False

BASE_DIR = Path(__file__).resolve().parent
OUT_DIR = BASE_DIR
OUT_DIR.mkdir(exist_ok=True)
os.chdir(BASE_DIR)
AUDIT_LOG = []
GLOBAL_DOC_MAP = {}

RU_TO_LAT = {
    'А': 'A', 'Б': 'B', 'В': 'V', 'Г': 'G', 'Д': 'D', 'Е': 'E', 'Ё': 'E',
    'Ж': 'ZH', 'З': 'Z', 'И': 'I', 'Й': 'I', 'К': 'K', 'Л': 'L', 'М': 'M',
    'Н': 'N', 'О': 'O', 'П': 'P', 'Р': 'R', 'С': 'S', 'Т': 'T', 'У': 'U',
    'Ф': 'F', 'Х': 'KH', 'Ц': 'TS', 'Ч': 'CH', 'Ш': 'SH', 'Щ': 'SHCH',
    'Ъ': '', 'Ы': 'Y', 'Ь': '', 'Э': 'E', 'Ю': 'IU', 'Я': 'IA'
}


def normalize_name_token(text) -> str:
    if text is None or pd.isna(text):
        return ""
    s = str(text).upper().strip()
    if s in ("N/A", "NONE", "NULL", "NOT PRESENTED", "NAN", ""):
        return ""
    s = re.sub(r"^(MR|MRS|MS|MISS|DR)\b\.?\s*", "", s)
    res = "".join(RU_TO_LAT.get(ch, ch) for ch in s)
    res = re.sub(r"[^A-Z]", "", res)
    res = res.replace("X", "KS").replace("TC", "TS")
    res = res.replace("YA", "IA").replace("YU", "IU").replace("YO", "E")
    res = res.replace("AY", "AI").replace("EY", "EI").replace("OY", "OI").replace("UY", "UI")
    res = re.sub(r"Y([AEIOU])", r"I\1", res)
    res = re.sub(r"(SKIY|SKII|SKY|SKI)$", "SKI", res)
    res = re.sub(r"(AYA|AIA)$", "AIA", res)
    res = re.sub(r"(IYA|IIA|IA)$", "IA", res)
    res = re.sub(r"(IY|YI|II|EI|EY|Y)$", "I", res)
    return res


def parse_single_date(val):
    if val is None or pd.isna(val):
        return None
    s = str(val).strip()
    if s in ("N/A", "None", "nan", "", "Not presented", "NULL"):
        return None
    try:
        if "/" in s:
            parts = s.split("/")
            if len(parts) == 3:
                if len(parts[0]) == 4:
                    y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                else:
                    m, d, y = int(parts[0]), int(parts[1]), int(parts[2])
                return f"{d:02d}.{m:02d}.{y:04d}"
        if "." in s:
            parts = s.split(".")
            if len(parts) == 3:
                if len(parts[0]) == 4:
                    y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                else:
                    d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
                return f"{d:02d}.{m:02d}.{y:04d}"
        if "-" in s:
            parts = s.split("-")
            if len(parts) == 3:
                if len(parts[0]) == 4:
                    y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                else:
                    d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
                return f"{d:02d}.{m:02d}.{y:04d}"
    except Exception:
        pass
    dt = pd.to_datetime(s, errors="coerce")
    return dt.strftime("%d.%m.%Y") if pd.notna(dt) else None


def normalize_date_series(s: pd.Series) -> pd.Series:
    return s.apply(parse_single_date)


def parse_raw_doc_digits(val):
    if val is None or pd.isna(val):
        return None
    s = str(val).upper().strip()
    if s in ("N/A", "NOT PRESENTED", "NONE", "NAN", "", "NULL"):
        return None
    m_int = re.match(r"^(\d{1,4})\s+(\d{6})$", s)
    if m_int:
        return m_int.group(1).zfill(4) + m_int.group(2)
    m_for = re.match(r"^(\d{1,2})\s+(\d{7})$", s)
    if m_for:
        return m_for.group(1).zfill(2) + m_for.group(2)
    digits = re.sub(r"[^0-9]", "", s)
    if not digits:
        return None
    return digits


def register_docs_in_global_map(series: pd.Series):
    for val in series.dropna():
        d = parse_raw_doc_digits(val)
        if not d:
            continue
        key = d.lstrip("0")
        if not key:
            continue
        prev = GLOBAL_DOC_MAP.get(key)
        if prev is None or len(d) > len(prev):
            GLOBAL_DOC_MAP[key] = d


def format_canonical_doc(val):
    d = parse_raw_doc_digits(val)
    if not d:
        return None
    key = d.lstrip("0")
    if key in GLOBAL_DOC_MAP and len(GLOBAL_DOC_MAP[key]) >= len(d):
        d = GLOBAL_DOC_MAP[key]
    if len(d) == 8:
        d = d.zfill(10)
    if len(d) == 10:
        return f"{d[:4]} {d[4:]}"
    if len(d) == 9:
        return f"{d[:2]} {d[2:]}"
    return d


def normalize_doc_series(s: pd.Series) -> pd.Series:
    return s.apply(format_canonical_doc)


def log_audit(dataset_name, source_count, parsed_count, unparsed_dates=0, notes=""):
    status = "OK" if source_count == parsed_count else f"DIFF: {parsed_count - source_count}"
    AUDIT_LOG.append({
        "Dataset": dataset_name,
        "Source_Rows": source_count,
        "Parquet_Rows": parsed_count,
        "Status": status,
        "Unparsed_Dates": unparsed_dates,
        "Notes": notes
    })


def pre_scan_documents():
    if os.path.exists("Sirena-export-fixed.tab"):
        with open("Sirena-export-fixed.tab", "r", encoding="utf-8", errors="replace") as f:
            lines = [line.rstrip("\r\n") for line in f if line.strip()]
        if lines:
            header = lines[0]
            if "\t" in header and len(header.split("\t")) >= 15:
                df_s = pd.read_csv("Sirena-export-fixed.tab", sep="\t", dtype=str, encoding="utf-8")
                df_s.columns = [c.strip() for c in df_s.columns]
                if "TravelDoc" in df_s.columns:
                    register_docs_in_global_map(df_s["TravelDoc"])
            else:
                docs = []
                row_re = re.compile(r"^.{108}[A-Z0-9]{4,6}(?:YES|NO)\s*[A-Z]{3}\s+[A-Z]{3}\s+[A-Z0-9]{6}[0-9]{16}\s+([0-9\s]{9,12})\s+")
                for line in lines[1:]:
                    m = row_re.match(line)
                    if m:
                        docs.append(m.group(1))
                register_docs_in_global_map(pd.Series(docs, dtype=str))

    if os.path.exists("BoardingData.csv"):
        with open("BoardingData.csv", "r", encoding="utf-8", errors="replace") as f:
            first_line = f.readline()
        sep = ";" if first_line.count(";") > first_line.count(",") else ","
        if first_line.count("\t") > max(first_line.count(";"), first_line.count(",")):
            sep = "\t"
        df_b = pd.read_csv("BoardingData.csv", sep=sep, dtype=str, encoding="utf-8", on_bad_lines="skip")
        if "PassengerDocument" in df_b.columns:
            register_docs_in_global_map(df_b["PassengerDocument"])


def parse_boarding_data(filepath="BoardingData.csv"):
    if not os.path.exists(filepath):
        return
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        first_line = f.readline()
        total_lines = sum(1 for line in f if line.strip())
    sep = ";" if first_line.count(";") > first_line.count(",") else ","
    if first_line.count("\t") > max(first_line.count(";"), first_line.count(",")):
        sep = "\t"

    df = pd.read_csv(filepath, sep=sep, dtype=str, encoding="utf-8", on_bad_lines="skip")
    df["norm_birth_date"] = normalize_date_series(df["PassengerBirthDate"])
    df["norm_flight_date"] = normalize_date_series(df["FlightDate"])

    raw_bd_valid = df["PassengerBirthDate"].dropna().replace({"Not presented": None, "N/A": None}).dropna()
    unparsed_bd = int(raw_bd_valid.notna().sum() - df["norm_birth_date"].notna().sum())

    df["norm_first_name"] = df["PassengerFirstName"].apply(normalize_name_token)
    df["norm_last_name"] = df["PassengerLastName"].apply(normalize_name_token)
    df["norm_second_name"] = df["PassengerSecondName"].apply(normalize_name_token)
    df["patronymic_initial"] = df["norm_second_name"].str.slice(0, 1)

    df["norm_doc"] = normalize_doc_series(df["PassengerDocument"])
    df["norm_ticket"] = df["TicketNumber"].replace({"Not presented": None, "N/A": None}).str.strip()
    df["norm_pnr"] = df["BookingCode"].replace({"Not presented": None, "N/A": None}).str.strip()

    df.to_parquet(OUT_DIR / "boarding_data.parquet", index=False)
    log_audit("BoardingData.csv", total_lines, len(df), unparsed_bd, f"sep='{sep}'")


def parse_sirena(filepath="Sirena-export-fixed.tab"):
    if not os.path.exists(filepath):
        return
    with open(filepath, "r", encoding="utf-8", errors="replace") as f:
        lines = [line.rstrip("\r\n") for line in f if line.strip()]

    header_line = lines[0]
    data_lines = lines[1:]

    if "\t" in header_line and len(header_line.split("\t")) >= 15:
        df = pd.read_csv(filepath, sep="\t", dtype=str, encoding="utf-8")
        df.columns = [c.strip() for c in df.columns]
        if "e-Ticket" in df.columns and "eTicket" not in df.columns:
            df["eTicket"] = df["e-Ticket"].str.strip()
        if "Fare" in df.columns:
            df["Fare"] = df["Fare"].str.strip()
    else:
        row_re = re.compile(
            r"^(?P<PaxName>.{60})"
            r"(?P<PaxBirthDate>.{12})"
            r"(?P<DepartDate>.{12})"
            r"(?P<DepartTime>.{12})"
            r"(?P<ArrivalDate>.{12})"
            r"(?P<ArrivalTime>.{12})"
            r"(?P<Flight>[A-Z0-9]{4,6})"
            r"(?P<CodeSh>YES|NO)\s*"
            r"(?P<From>[A-Z]{3})\s+"
            r"(?P<Dest>[A-Z]{3})\s+"
            r"(?P<Code>[A-Z0-9]{6})"
            r"(?P<eTicket>[0-9]{16})\s+"
            r"(?P<TravelDoc>[0-9\s]{9,12})\s+"
            r"(?P<Rest>.*)$"
        )
        parsed_rows = []
        for line in data_lines:
            m = row_re.match(line)
            if m:
                d = {k: v.strip() for k, v in m.groupdict().items()}
                rest = d.pop("Rest", "")
                ff_match = re.search(r"(FF#[A-Z]{2}\s*\d+)", rest)
                fare_match = re.search(r"\b([A-Z][A-Z0-9]{5})\b", rest)
                d["PaxAdditionalInfo"] = ff_match.group(1) if ff_match else ""
                d["Fare"] = fare_match.group(1) if fare_match else ""
                d["RawRest"] = rest
                parsed_rows.append(d)
            else:
                parsed_rows.append({"PaxName": line.strip(), "Fare": "", "RawRest": "UNPARSED"})
        df = pd.DataFrame(parsed_rows)

    fio_parts = df["PaxName"].astype(str).str.strip().str.split(r"\s+", n=2, expand=True)
    df["raw_last_ru"] = fio_parts[0] if 0 in fio_parts else ""
    df["raw_first_ru"] = fio_parts[1] if 1 in fio_parts else ""
    df["raw_second_ru"] = fio_parts[2] if 2 in fio_parts else ""

    df["norm_last_name"] = df["raw_last_ru"].apply(normalize_name_token)
    df["norm_first_name"] = df["raw_first_ru"].apply(normalize_name_token)
    df["norm_second_name"] = df["raw_second_ru"].apply(normalize_name_token)
    df["patronymic_initial"] = df["norm_second_name"].str.slice(0, 1)

    df["norm_birth_date"] = normalize_date_series(df.get("PaxBirthDate", pd.Series(dtype=str)))
    df["norm_flight_date"] = normalize_date_series(df.get("DepartDate", pd.Series(dtype=str)))
    df["norm_arrival_date"] = normalize_date_series(df.get("ArrivalDate", pd.Series(dtype=str)))
    df["norm_doc"] = normalize_doc_series(df.get("TravelDoc", pd.Series(dtype=str)))

    if "PaxAdditionalInfo" in df.columns:
        df["norm_ff_card"] = (
            df["PaxAdditionalInfo"]
            .astype(str)
            .str.extract(r"FF#([A-Z]{2}\s*\d+)", expand=False)
            .str.replace(r"\s+", " ", regex=True)
            .str.strip()
        )

    df.to_parquet(OUT_DIR / "sirena.parquet", index=False)
    log_audit("Sirena-export-fixed.tab", len(data_lines), len(df), 0, "normalized")


def extract_bp_from_grid(grid, filename, sheetname):
    def cell(r, c):
        if r < len(grid) and c < len(grid[r]):
            val = grid[r][c]
            return str(val).strip() if val is not None and str(val).strip() != "" else None
        return None

    seq_num = cell(0, 7) or cell(0, 6)
    title = cell(2, 0)
    full_name = cell(2, 1) or cell(2, 2)
    travel_class = cell(2, 7)
    flight_num = cell(4, 0)
    from_city = cell(4, 3)
    to_city = cell(4, 7)
    gate = cell(6, 1)
    from_iata = cell(6, 3)
    to_iata = cell(6, 7)
    raw_date = cell(8, 0)
    raw_time = cell(8, 2) or cell(8, 1)
    operator = cell(8, 4)
    seat = cell(10, 7)
    pnr = cell(12, 1)
    eticket = cell(12, 4) if cell(12, 3) == "E-TICKET" else (cell(12, 4) or cell(12, 3))

    name_parts = (full_name or "").split()
    p1 = normalize_name_token(name_parts[0]) if len(name_parts) >= 1 else ""
    p2 = normalize_name_token(name_parts[-1]) if len(name_parts) >= 2 else ""

    return {
        "source_file": filename,
        "sheet_name": sheetname,
        "sequence": seq_num,
        "title": title,
        "raw_name": full_name,
        "norm_name_part1": p1,
        "norm_name_part2": p2,
        "travel_class": travel_class,
        "flight_number": flight_num,
        "from_city": from_city,
        "to_city": to_city,
        "from_iata": from_iata.upper() if from_iata else None,
        "to_iata": to_iata.upper() if to_iata else None,
        "gate": gate,
        "raw_date": raw_date,
        "flight_time": raw_time,
        "operator": operator,
        "seat": seat,
        "pnr": pnr,
        "e_ticket": eticket
    }


def parse_boarding_passes_folder(folder="YourBoardingPassDotAero"):
    if not os.path.exists(folder):
        return
    xlsx_files = sorted(glob.glob(os.path.join(folder, "*.xlsx")))
    all_passes = []
    total_sheets = 0

    for fpath in xlsx_files:
        fname = os.path.basename(fpath)
        if HAS_CALAMINE and CalamineWorkbook is not None:
            wb = CalamineWorkbook.from_path(fpath)
            for sname in wb.sheet_names:
                total_sheets += 1
                grid = wb.get_sheet_by_name(sname).to_python()
                if grid:
                    all_passes.append(extract_bp_from_grid(grid, fname, sname))
        else:
            wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
            for sname in wb.sheetnames:
                total_sheets += 1
                ws = wb[sname]
                grid = [[cell for cell in row] for row in ws.iter_rows(max_row=15, max_col=10, values_only=True)]
                if grid:
                    all_passes.append(extract_bp_from_grid(grid, fname, sname))
            wb.close()

    df = pd.DataFrame(all_passes)
    file_date = df["source_file"].str.extract(r"(\d{4}-\d{2}-\d{2})", expand=False)
    cell_date = df["raw_date"].astype(str).str.extract(r"(\d{4}-\d{2}-\d{2})", expand=False)
    combined_date = cell_date.fillna(file_date)
    df["norm_flight_date"] = normalize_date_series(combined_date)

    df.to_parquet(OUT_DIR / "boarding_passes.parquet", index=False)
    log_audit("YourBoardingPassDotAero", total_sheets, len(df), int(df["norm_flight_date"].isna().sum()), f"files={len(xlsx_files)}")


def parse_pointz_xml(filepath="PointzAggregator-AirlinesData.xml"):
    if not os.path.exists(filepath):
        if os.path.exists("PointzAggregator-AirlinesData"):
            filepath = "PointzAggregator-AirlinesData"
        else:
            return

    activities = []
    cards_only = []
    raw_activity_count = 0
    raw_user_count = 0

    for _, elem in ET.iterparse(filepath, events=("end",)):
        if elem.tag == "user":
            raw_user_count += 1
            uid = elem.get("uid")
            name_el = elem.find("name")
            first = name_el.get("first", "") if name_el is not None else ""
            last = name_el.get("last", "") if name_el is not None else ""
            norm_first = normalize_name_token(first)
            norm_last = normalize_name_token(last)

            for card in elem.findall(".//card"):
                c_num = re.sub(r"\s+", " ", (card.get("number") or "").strip())
                prog = (card.findtext("bonusprogramm") or "").strip()
                acts = card.findall(".//activity")

                cards_only.append({
                    "uid": uid,
                    "raw_first_name": first,
                    "raw_last_name": last,
                    "norm_first_name": norm_first,
                    "norm_last_name": norm_last,
                    "card_number": c_num,
                    "bonus_program": prog,
                    "activities_count": len(acts)
                })

                for act in acts:
                    raw_activity_count += 1
                    activities.append({
                        "uid": uid,
                        "raw_first_name": first,
                        "raw_last_name": last,
                        "norm_first_name": norm_first,
                        "norm_last_name": norm_last,
                        "card_number": c_num,
                        "bonus_program": prog,
                        "flight_number": (act.findtext("Code") or "").strip(),
                        "flight_date": parse_single_date(act.findtext("Date")),
                        "origin": (act.findtext("Departure") or "").strip().upper(),
                        "destination": (act.findtext("Arrival") or "").strip().upper(),
                        "fare": (act.findtext("Fare") or "").strip()
                    })
            elem.clear()

    pd.DataFrame(activities).to_parquet(OUT_DIR / "pointz_activities.parquet", index=False)
    pd.DataFrame(cards_only).to_parquet(OUT_DIR / "pointz_users_cards.parquet", index=False)
    log_audit("PointzAggregator.xml", raw_activity_count, len(activities), 0, f"users={raw_user_count}")


def parse_forum_json(filepath="FrequentFlyerForum-Profiles.json"):
    if not os.path.exists(filepath):
        return
    with open(filepath, "r", encoding="utf-8") as f:
        data = json.load(f)

    profiles = data.get("Forum Profiles", [])
    users_list = []
    flights_list = []
    loyalty_list = []
    raw_flights_count = 0

    for idx, prof in enumerate(profiles):
        nick = prof.get("NickName")
        sex = prof.get("Sex")
        dob = parse_single_date(prof.get("Date of Birth"))
        real_name = prof.get("Real Name") or {}
        r_first = real_name.get("First Name")
        r_last = real_name.get("Last Name")

        passports = []
        for doc_obj in prof.get("Travel Documents", []):
            if isinstance(doc_obj, dict) and doc_obj.get("Passports"):
                fmt_p = format_canonical_doc(doc_obj.get("Passports"))
                if fmt_p:
                    passports.append(fmt_p)

        users_list.append({
            "profile_id": idx,
            "nickname": nick,
            "sex": sex,
            "birth_date": dob,
            "raw_first_name": r_first,
            "raw_last_name": r_last,
            "norm_first_name": normalize_name_token(r_first),
            "norm_last_name": normalize_name_token(r_last),
            "passports": ", ".join(passports)
        })

        for lp in prof.get("Loyality Programm", []):
            if isinstance(lp, dict):
                prog_code = str(lp.get("programm") or "").strip()
                num = str(lp.get("Number") or "").strip()
                loyalty_list.append({
                    "profile_id": idx,
                    "nickname": nick,
                    "program": prog_code,
                    "number": num,
                    "norm_ff_card": f"{prog_code} {num}".strip(),
                    "status": lp.get("Status")
                })

        for fl in prof.get("Registered Flights", []):
            if isinstance(fl, dict):
                raw_flights_count += 1
                dep = fl.get("Departure") or {}
                arr = fl.get("Arrival") or {}
                flights_list.append({
                    "profile_id": idx,
                    "nickname": nick,
                    "norm_first_name": normalize_name_token(r_first),
                    "norm_last_name": normalize_name_token(r_last),
                    "flight_date": parse_single_date(fl.get("Date")),
                    "flight_number": fl.get("Flight"),
                    "codeshare": fl.get("Codeshare"),
                    "from_iata": str(dep.get("Airport") or "").upper(),
                    "from_city": dep.get("City"),
                    "to_iata": str(arr.get("Airport") or "").upper(),
                    "to_city": arr.get("City")
                })

    pd.DataFrame(users_list).to_parquet(OUT_DIR / "forum_users.parquet", index=False)
    pd.DataFrame(loyalty_list).to_parquet(OUT_DIR / "forum_loyalty.parquet", index=False)
    pd.DataFrame(flights_list).to_parquet(OUT_DIR / "forum_flights.parquet", index=False)
    log_audit("FrequentFlyerForum.json", raw_flights_count, len(flights_list), 0, f"profiles={len(users_list)}")


def parse_skyteam_yaml(filepath="SkyTeam-Exchange.yaml"):
    if not os.path.exists(filepath):
        return
    with open(filepath, "r", encoding="utf-8") as f:
        data = yaml.load(f, Loader=YamlLoader)

    rows = []
    raw_count = 0
    if isinstance(data, dict):
        for f_date, flights in data.items():
            if not isinstance(flights, dict):
                continue
            clean_date = parse_single_date(f_date)
            for f_code, info in flights.items():
                if not isinstance(info, dict):
                    continue
                origin = str(info.get("FROM", "")).upper()
                dest = str(info.get("TO", "")).upper()
                status = info.get("STATUS")
                ff_dict = info.get("FF") or {}
                if isinstance(ff_dict, dict):
                    for card_num, t_info in ff_dict.items():
                        raw_count += 1
                        rows.append({
                            "flight_date": clean_date,
                            "flight_number": str(f_code),
                            "from_iata": origin,
                            "to_iata": dest,
                            "status": status,
                            "norm_ff_card": re.sub(r"\s+", " ", str(card_num).strip()),
                            "travel_class": t_info.get("CLASS") if isinstance(t_info, dict) else None,
                            "fare": t_info.get("FARE") if isinstance(t_info, dict) else None
                        })

    pd.DataFrame(rows).to_parquet(OUT_DIR / "skyteam_exchange.parquet", index=False)
    log_audit("SkyTeam-Exchange.yaml", raw_count, len(rows), 0, "normalized")


def parse_timetable_pdf():
    pdf_candidates = glob.glob("Skyteam_Timetable*.pdf")
    if not pdf_candidates:
        return
    filepath = pdf_candidates[0]
    from pypdf import PdfReader
    reader = PdfReader(filepath)
    pages_data = []
    for i, page in enumerate(reader.pages):
        txt = page.extract_text() or ""
        pages_data.append({"page_number": i + 1, "raw_text": txt})

    pd.DataFrame(pages_data).to_parquet(OUT_DIR / "timetable_pages.parquet", index=False)
    log_audit("Skyteam_Timetable.pdf", len(reader.pages), len(pages_data), 0, "pages")


def build_duckdb_and_report():
    db_path = BASE_DIR / "yoyoflot.duckdb"
    con = duckdb.connect(db_path)
    use_existing = not AUDIT_LOG
    for pq_file in OUT_DIR.glob("*.parquet"):
        table_name = pq_file.stem
        con.execute(f"CREATE OR REPLACE VIEW {table_name} AS SELECT * FROM '{pq_file.as_posix()}'")
        if use_existing:
            row_count = con.execute(f"SELECT COUNT(*) FROM '{pq_file.as_posix()}'").fetchone()[0]
            log_audit(pq_file.name, row_count, row_count, 0, "existing parquet")
    con.execute("CREATE SCHEMA IF NOT EXISTS source")
    for pq_file in OUT_DIR.glob("*.parquet"):
        table_name = pq_file.stem.replace('"', '""')
        con.execute(f'CREATE OR REPLACE VIEW source."{table_name}" AS SELECT * FROM main."{table_name}"')
    con.execute("""
        CREATE OR REPLACE VIEW source.sirena AS
        SELECT *, eTicket AS "e-Ticket" FROM main.sirena;

        CREATE OR REPLACE VIEW source.pointz_activities AS
        SELECT a.*, c.activities_count
        FROM main.pointz_activities a
        LEFT JOIN main.pointz_users_cards c
          ON a.uid = c.uid AND a.card_number = c.card_number;

        CREATE OR REPLACE VIEW source."boarding_data + boarding_passes" AS
        WITH matches AS (
            SELECT bp.record_id AS pass_id, b.record_id AS boarding_id
            FROM main.boarding_passes bp
            JOIN main.sirena s ON bp.e_ticket = s.eTicket
            JOIN main.boarding_data b
              ON b.norm_doc = s.norm_doc
             AND UPPER(TRIM(b.FlightNumber)) = UPPER(TRIM(s.Flight))
             AND b.norm_flight_date = s.norm_flight_date
            QUALIFY ROW_NUMBER() OVER (
                PARTITION BY bp.record_id ORDER BY b.source_row
            ) = 1
        ), joined AS (
            SELECT bp.*, b.*,
                   bp.record_id AS pass_record_id,
                   b.record_id AS boarding_record_id
            FROM main.boarding_passes bp
            FULL OUTER JOIN matches m ON bp.record_id = m.pass_id
            FULL OUTER JOIN main.boarding_data b ON b.record_id = m.boarding_id
        )
        SELECT * EXCLUDE (pass_record_id, boarding_record_id) FROM joined;
    """)
    con.close()

    audit_df = pd.DataFrame(AUDIT_LOG)
    with open(BASE_DIR / "etl_audit_report.txt", "w", encoding="utf-8") as f:
        f.write("ETL DATA INTEGRITY REPORT\n\n")
        f.write(audit_df.to_string(index=False) + "\n")


if __name__ == "__main__":
    pre_scan_documents()
    parse_boarding_data()
    parse_sirena()
    parse_boarding_passes_folder()
    parse_pointz_xml()
    parse_forum_json()
    parse_skyteam_yaml()
    parse_timetable_pdf()
    build_duckdb_and_report()
